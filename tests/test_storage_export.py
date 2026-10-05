import io
from datetime import date

import pytest
from openpyxl import load_workbook

from receipts.export import export_excel
from receipts.models import LineItem, Receipt, receipt_from_dict
from receipts.storage import (
    DuplicateReceipt, connect, delete_receipt, file_hash, items_df, receipts_df, save_receipt,
)


def _receipt(toko="Toko A", total=60000):
    return Receipt(toko=toko, tanggal=date(2026, 9, 12), items=[LineItem("Beras", 2, 30000, 60000)],
                   total=total, engine="test")


def test_save_and_read_back():
    conn = connect()
    rid = save_receipt(conn, _receipt(), file_hash=file_hash(b"a"), filename="a.jpg")
    assert rid == 1
    rec, items = receipts_df(conn), items_df(conn)
    assert rec.loc[0, "toko"] == "Toko A" and rec.loc[0, "total"] == 60000
    assert items.loc[0, "nama"] == "Beras" and items.loc[0, "nota_id"] == rid


def test_duplicate_photo_is_rejected():
    conn = connect()
    save_receipt(conn, _receipt(), file_hash=file_hash(b"a"))
    with pytest.raises(DuplicateReceipt):
        save_receipt(conn, _receipt(), file_hash=file_hash(b"a"))
    assert len(receipts_df(conn)) == 1  # tidak ada nota setengah tersimpan


def test_delete_cascades_to_items():
    conn = connect()
    rid = save_receipt(conn, _receipt(), file_hash=file_hash(b"a"))
    delete_receipt(conn, rid)
    assert receipts_df(conn).empty and items_df(conn).empty


def test_export_excel_sheets_and_sum_formula():
    conn = connect()
    save_receipt(conn, _receipt("Toko A", 60000), file_hash=file_hash(b"a"))
    save_receipt(conn, _receipt("Toko B", 40000), file_hash=file_hash(b"b"))
    buf = io.BytesIO()
    export_excel(receipts_df(conn), items_df(conn), buf)
    wb = load_workbook(io.BytesIO(buf.getvalue()))
    assert wb.sheetnames == ["Rekap nota", "Detail item", "Per toko", "Per bulan"]
    ws = wb["Rekap nota"]
    assert ws.cell(1, 8).value == "Total"
    assert str(ws.cell(4, 8).value).startswith("=SUM(")  # baris total memakai rumus
    assert wb["Per toko"].max_row == 4  # judul + 2 toko + baris TOTAL


def test_receipt_from_dict_validates_types():
    r = receipt_from_dict({"toko": "X", "tanggal": "2026-09-12", "total": "60.000",
                           "items": [{"nama": "Beras", "qty": 2, "harga_satuan": 30000},
                                     {"nama": "", "subtotal": 5}, "sampah"]})
    assert r.total == 60000 and r.tanggal == date(2026, 9, 12)
    assert len(r.items) == 1 and r.items[0].subtotal == 60000  # subtotal dihitung dari qty x harga
    assert receipt_from_dict({"tanggal": "bukan-tanggal"}).tanggal is None
