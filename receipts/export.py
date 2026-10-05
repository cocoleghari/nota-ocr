"""Ekspor rekap ke Excel dengan beberapa lembar dan format rapi."""
from __future__ import annotations

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="3B5BDB")
HEADER_FONT = Font(bold=True, color="FFFFFF")
MONEY_COLS = {"subtotal", "pajak", "diskon", "total", "harga_satuan", "total_belanja", "rata_rata"}
LABELS = {
    "id": "No. nota", "tanggal": "Tanggal", "waktu": "Waktu", "toko": "Toko", "subtotal": "Subtotal",
    "pajak": "Pajak & layanan", "diskon": "Diskon", "total": "Total", "nama": "Item", "qty": "Qty",
    "harga_satuan": "Harga satuan", "jumlah_nota": "Jumlah nota", "total_belanja": "Total belanja",
    "rata_rata": "Rata-rata per nota", "bulan": "Bulan", "nota_id": "No. nota",
}


def _summaries(receipts: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    per_toko = (receipts.groupby("toko", dropna=False)["total"]
                .agg(jumlah_nota="count", total_belanja="sum", rata_rata="mean")
                .round(0).reset_index().sort_values("total_belanja", ascending=False))
    per_toko["toko"] = per_toko["toko"].fillna("(tanpa nama)")
    by_month = receipts.dropna(subset=["tanggal"]).assign(bulan=lambda d: d["tanggal"].dt.strftime("%Y-%m"))
    per_bulan = (by_month.groupby("bulan")["total"].agg(jumlah_nota="count", total_belanja="sum")
                 .reset_index().sort_values("bulan"))
    return per_toko, per_bulan


def _style(ws, df: pd.DataFrame, total_col: str | None = None) -> None:
    cols = list(df.columns)
    for c in ws[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(vertical="center")
    for idx, name in enumerate(cols, start=1):
        letter = get_column_letter(idx)
        width = max([len(str(ws.cell(1, idx).value or ""))] + [len(str(v)) for v in df[name].head(200)]) + 3
        ws.column_dimensions[letter].width = min(max(width, 10), 42)
        for row in range(2, len(df) + 2):
            cell = ws.cell(row, idx)
            if name in MONEY_COLS:
                cell.number_format = "#,##0"
            elif name == "tanggal":
                cell.number_format = "DD MMM YYYY"
    ws.freeze_panes = "A2"
    if total_col and total_col in cols and len(df):
        idx = cols.index(total_col) + 1
        letter = get_column_letter(idx)
        last = len(df) + 1
        ws.cell(last + 1, idx - 1 if idx > 1 else idx, "TOTAL").font = Font(bold=True)
        cell = ws.cell(last + 1, idx, f"=SUM({letter}2:{letter}{last})")
        cell.font, cell.number_format = Font(bold=True), "#,##0"


def export_excel(receipts: pd.DataFrame, items: pd.DataFrame, target) -> None:
    """Tulis rekap ke `target` (path atau buffer). Baris TOTAL memakai rumus SUM agar tetap hidup di Excel."""
    rekap = receipts[["id", "tanggal", "waktu", "toko", "subtotal", "pajak", "diskon", "total"]].copy()
    detail = items[["nota_id", "tanggal", "toko", "nama", "qty", "harga_satuan", "subtotal"]].copy()
    per_toko, per_bulan = _summaries(receipts)
    sheets = [("Rekap nota", rekap, "total"), ("Detail item", detail, "subtotal"),
              ("Per toko", per_toko, "total_belanja"), ("Per bulan", per_bulan, "total_belanja")]
    with pd.ExcelWriter(target, engine="openpyxl", datetime_format="DD MMM YYYY") as xw:
        for name, df, total_col in sheets:
            out = df.rename(columns=LABELS)
            out.to_excel(xw, sheet_name=name, index=False)
            _style(xw.sheets[name], df, total_col)
            # nama kolom asli dipakai untuk format; judul tampil dengan label Indonesia
